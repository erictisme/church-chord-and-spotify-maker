#!/usr/bin/env python3
"""Assemble a Sunday service chord doc from a week config.

Reads a JSON config (see weeks/), pulls each song from the library,
transposes to the requested key with column-preserving logic, and writes:
  - output/<display_date>.docx  — two columns, Times New Roman, ready for Drive
  - output/<display_date>.txt   — plain-text fallback for copy-paste

Usage:
    python3 build_service.py weeks/2026-06-14.json
"""

import json
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Pt
from PIL import ImageFont

from transpose import is_chord_line, parse_song_file, transpose_chart

HERE = os.path.dirname(os.path.abspath(__file__))
LIBRARY = os.path.join(HERE, "..", "eric-songs", "all-songs")
OUTPUT = os.path.join(HERE, "output")

FONT = "Times New Roman"
FONT_SIZE = 11
FONT_FILE = "/System/Library/Fonts/Supplemental/Times New Roman.ttf"

# Times New Roman is proportional: a space is ~half a letter wide, so
# space-aligned chords drift left in Word. Instead, measure where each chord's
# column lands in the lyric below (real glyph widths) and place it with a tab stop.
_SCALE = 100
_measure_font = ImageFont.truetype(FONT_FILE, FONT_SIZE * _SCALE,
                                   layout_engine=ImageFont.Layout.BASIC)


def text_width_pt(text):
    return _measure_font.getlength(text) / _SCALE


def chord_positions(chord_line, lyric):
    """Return [(x_pt, chord)] aligning each chord over its lyric column."""
    space_w = text_width_pt(" ")
    out, min_x = [], 0.0
    for m in re.finditer(r"\S+", chord_line):
        col, chord = m.start(), m.group()
        if col <= len(lyric):
            x = text_width_pt(lyric[:col])
        else:  # chord hangs past the end of the lyric
            x = text_width_pt(lyric) + (col - len(lyric)) * space_w * 2
        x = min(x, COLUMN_W - text_width_pt(chord) - 2)  # stay inside the column
        x = max(x, min_x)  # never let chords collide
        out.append((x, chord))
        min_x = x + text_width_pt(chord + " ")
    return out


def load_slot_chart(slot):
    """Return (title, key, chart_text) for a slot, transposed if needed."""
    meta, chart = parse_song_file(os.path.join(LIBRARY, slot["file"]))
    source_key = meta.get("key", "C")
    target_key = slot.get("key") or source_key
    if target_key != source_key:
        chart = transpose_chart(chart, source_key, target_key)
    return meta.get("title", slot["file"]), target_key, chart


def build_header_lines(cfg, slot_titles):
    lines = [cfg["display_date"],
             f"Theme: {cfg['theme']}",
             f"Scripture: {cfg['scripture']}",
             ""]
    current_section = None
    for slot, title in slot_titles:
        if slot["section"] != current_section:
            current_section = slot["section"]
            lines.append(f"{current_section}:")
        entry = title if title else slot.get("note", "TBC")
        if title and slot.get("note"):
            entry += f" - {slot['note']}"
        lines.append(entry)
    return lines


def two_columns(section):
    """Switch a docx section to a two-column layout."""
    cols = section._sectPr.xpath("./w:cols")[0]
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "360")  # 0.25" gutter


def clean_line(line):
    """Strip markdown bold markers that appear in some library files."""
    return line.replace("**", "")


def clean_chart(chart):
    """Normalise markdown-style library files (## headers, ``` fences,
    **Key:**/**Artist:** lines) to the plain [Section] format."""
    out = []
    for line in chart.splitlines():
        s = line.strip()
        if s.startswith("```"):
            continue
        if re.match(r"^\*\*(Key|Artist):", s):
            continue
        if s.startswith("## "):
            line = f"[{s[3:].strip()}]"
        out.append(line)
    text = "\n".join(out)
    return re.sub(r"\n{3,}", "\n\n", text).strip("\n")


def split_to_fit(chord_line, lyric, max_w):
    """Split a chord+lyric pair at a word break so the lyric fits one column.
    Returns [(chord_line, lyric), ...]; chords stay over their words."""
    if text_width_pt(lyric) <= max_w:
        return [(chord_line, lyric)]
    cut = None
    for m in re.finditer(r" +", lyric):
        if text_width_pt(lyric[:m.start()]) > max_w:
            break
        cut = m
    if cut is None:
        return [(chord_line, lyric)]
    a, b = cut.start(), cut.end()
    head = [(m.start(), m.group()) for m in re.finditer(r"\S+", chord_line) if m.start() < a]
    tail = [(max(m.start() - b, 0), m.group()) for m in re.finditer(r"\S+", chord_line) if m.start() >= a]
    def render(tokens):
        line = ""
        for col, tok in tokens:
            line = line.ljust(col) if len(line) < col else (line + " " if line else line)
            line += tok
        return line
    return [(render(head), lyric[:a])] + split_to_fit(render(tail), lyric[b:], max_w)


def is_section_header(line):
    """True for [Verse 1], [Chorus], [Bridge] etc."""
    s = line.strip()
    return s.startswith("[") and s.endswith("]")


def add_chord_paragraph(doc, chord_line, lyric):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1
    parts = []
    for x, chord in chord_positions(chord_line, lyric):
        if x > 0.05:
            p.paragraph_format.tab_stops.add_tab_stop(Pt(x), WD_TAB_ALIGNMENT.LEFT)
            parts.append("\t" + chord)
        else:
            parts.append(chord)
    p.add_run("".join(parts))


def is_lyric_line(line):
    s = line.strip()
    return bool(s) and not is_section_header(line) and not is_chord_line(line)


COLUMN_W = 225.0  # pt; set from the real page in build()


def add_lines(doc, lines, bold=False):
    lines = [clean_line(l) for l in lines]
    skip = False
    for i, line in enumerate(lines):
        if skip:
            skip = False
            continue
        nxt = lines[i + 1] if i + 1 < len(lines) else ""
        if not bold and is_chord_line(line) and is_lyric_line(nxt):
            for ch, ly in split_to_fit(line, nxt, COLUMN_W - 4):
                if ch.strip():
                    add_chord_paragraph(doc, ch, ly)
                add_lines(doc, [ly])
            skip = True
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1
        if is_section_header(line):
            p.paragraph_format.space_before = Pt(6)
        else:
            p.paragraph_format.space_before = Pt(0)
        run = p.add_run(line)
        run.bold = bold
    return doc


def build(cfg_path):
    cfg = json.load(open(cfg_path, encoding="utf-8"))
    os.makedirs(OUTPUT, exist_ok=True)

    # Resolve every slot up front so header and body agree.
    resolved = []  # (slot, title, key, chart) — title None for TBC slots
    for slot in cfg["slots"]:
        if slot.get("file"):
            title, key, chart = load_slot_chart(slot)
            resolved.append((slot, title, key, chart))
        else:
            resolved.append((slot, None, None, None))

    header = build_header_lines(cfg, [(s, t) for s, t, _, _ in resolved])

    # --- plain text ---
    txt_parts = ["\n".join(header)]
    for slot, title, key, chart in resolved:
        if title:
            txt_parts.append(f"\n\n{title} ({key})\n{clean_line(clean_chart(chart))}")
    txt = "".join(txt_parts) + "\n"

    # --- docx ---
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(FONT_SIZE)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    two_columns(doc.sections[0])
    global COLUMN_W
    sec = doc.sections[0]
    COLUMN_W = (sec.page_width - sec.left_margin - sec.right_margin) / 12700 / 2 - 9

    add_lines(doc, header[:1], bold=True)      # date line bold
    add_lines(doc, header[1:])
    for slot, title, key, chart in resolved:
        if not title:
            continue
        add_lines(doc, [""])
        if slot.get("page_break"):
            doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
        add_lines(doc, [f"{title} ({key})"], bold=True)
        add_lines(doc, clean_chart(chart).splitlines())

    base = os.path.join(OUTPUT, cfg["display_date"])
    with open(base + ".txt", "w", encoding="utf-8") as f:
        f.write(txt)
    doc.save(base + ".docx")
    print(f"Wrote: {base}.docx")
    print(f"Wrote: {base}.txt")
    transposed = [f"{t} -> {k}" for s, t, k, _ in resolved
                  if t and s.get("key") and s["key"] != parse_song_file(
                      os.path.join(LIBRARY, s["file"]))[0].get("key", "C")]
    if transposed:
        print("Transposed: " + "; ".join(transposed))


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    build(sys.argv[1])
